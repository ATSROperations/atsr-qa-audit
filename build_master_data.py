"""
Rebuild js/master-data.js from the ATSR Master File.

Usage:  python3 tools/build_master_data.py "ATSR_Master_File.xlsx"
Needs:  pip install openpyxl

Reads the Data sheet (RTO table in columns A-E, course table in columns G-L)
and the Archived RTO sheet. College process notes are NOT here; they live in
js/config.js under COLLEGE_NOTES so a rebuild never wipes them.
"""
import json, sys, datetime
import openpyxl

src = sys.argv[1] if len(sys.argv) > 1 else "ATSR_Master_File.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Data"]

def clean(v):
    return str(v).strip() if v is not None else ""

rtos, quals, seen = [], [], set()
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0]:
        code = clean(r[0])
        if code not in seen:
            seen.add(code)
            rtos.append({"code": code, "name": clean(r[1]), "level": clean(r[2]),
                         "type": clean(r[3]), "processTo": clean(r[4]), "archived": False})
    if r[6]:
        quals.append({"code": clean(r[6]), "title": clean(r[7]), "package": clean(r[9]), "entry": clean(r[11])})

if "Archived RTO" in wb.sheetnames:
    for r in wb["Archived RTO"].iter_rows(min_row=2, values_only=True):
        if r[0] and clean(r[0]) not in seen:
            seen.add(clean(r[0]))
            rtos.append({"code": clean(r[0]), "name": clean(r[1]), "level": clean(r[2]),
                         "type": "", "processTo": "", "archived": True})

rtos.sort(key=lambda x: x["name"].lower())
quals.sort(key=lambda x: x["code"])

out = ("/* Generated from the ATSR Master File on " + datetime.date.today().isoformat() + ".\n"
       "   Rebuild with tools/build_master_data.py, or edit rows here directly.\n"
       "   level: Compliant | Less Compliant | Non Compliant. archived: no longer used for new files. */\n\n"
       "const MASTER = {\n  rtos: [\n" + ",\n".join("    " + json.dumps(x, ensure_ascii=False) for x in rtos) + "\n  ],\n"
       "  qualifications: [\n" + ",\n".join("    " + json.dumps(x, ensure_ascii=False) for x in quals) + "\n  ],\n};\n")
open("js/master-data.js", "w", encoding="utf-8").write(out)
print(len(rtos), "RTOs,", len(quals), "qualifications written to js/master-data.js")
